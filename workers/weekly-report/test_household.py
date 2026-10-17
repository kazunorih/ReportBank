import io
import unittest
from household import parse_income, parse_unemployment, collect_household


class HouseholdTests(unittest.TestCase):
    def test_income_reference_year_and_units(self):
        row = parse_income('平均所得金額（575万2千円）以下の割合は61.5％'.encode(), 2025)
        self.assertEqual((row['year'], row['averageIncome'], row['belowMeanShare']), (2024, 575.2, 61.5))
        row = parse_income('平均額（659万6千円）より低い世帯の割合は60.0％'.encode('cp932'), 1996)
        self.assertEqual(row['comparison'], 'below')
        with self.assertRaises(ValueError):
            parse_income('高齢者世帯の平均所得は336万円'.encode(), 2025)

    def test_annual_unemployment_uses_both_sexes_and_percentage_not_count(self):
        import openpyxl
        book = openpyxl.Workbook(); sheet = book.active
        sheet.cell(1, 3, '就業状態別15歳以上人口－全国')
        sheet.cell(5, 3, '男女計'); sheet.cell(6, 13, '完全失業率（％）')
        sheet.cell(11, 2, 2011); sheet.cell(11, 6, 999); sheet.cell(11, 13, '<4.6>')
        raw = io.BytesIO(); book.save(raw)
        self.assertEqual(parse_unemployment(raw.getvalue()), [{'year': 2011, 'unemploymentRate': 4.6}])
        sheet.cell(5, 3, '男'); raw = io.BytesIO(); book.save(raw)
        with self.assertRaises(ValueError):
            parse_unemployment(raw.getvalue())

    def test_cancelled_survey_not_filled_and_unpublished_year_not_invented(self):
        data = collect_household(None, '2026-10-17')
        rows = {r['year']: r for r in data['observations']}
        self.assertIsNone(rows[2019]['averageIncome'])
        self.assertIsNone(rows[2025]['averageIncome'])
        self.assertEqual(rows[2025]['unemploymentRate'], 2.5)
        self.assertNotIn(2026, rows)


if __name__ == '__main__':
    unittest.main()
